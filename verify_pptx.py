import sys
from pptx import Presentation

# Ensure stdout handles unicode without crashing
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

prs = Presentation('Event_Management_System_FYP_Presentation.pptx')
print(f'Total Slides: {len(prs.slides)}')
for idx, slide in enumerate(prs.slides):
    texts = []
    images = []
    tables = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    texts.append(p.text.strip())
        elif shape.has_table:
            tables.append(f'Table({len(shape.table.rows)}x{len(shape.table.columns)})')
        elif shape.shape_type == 13: # Picture
            images.append('Picture')
    first_title = texts[1] if len(texts) > 1 else (texts[0] if texts else 'No text')
    print(f'Slide {idx+1:02d}: shapes={len(slide.shapes)}, images={len(images)}, tables={len(tables)}, title="{first_title[:45]}"')
